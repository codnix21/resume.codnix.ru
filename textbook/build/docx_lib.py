# -*- coding: utf-8 -*-
"""Сборка рукописи: стили Word, разбор простой разметки, таблицы и листинги."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path("/workspace")
STAND = ROOT / "textbook" / "stand"
FONT = "Times New Roman"
MONO = "Liberation Mono"


def _set_run_font(run, name, size, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _style_font(style, name, size, bold=False, italic=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _paragraph_format(pf, *, before=0, after=6, line=1.5, first=None, left=None, align="justify"):
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if first is not None:
        pf.first_line_indent = Cm(first)
    else:
        pf.first_line_indent = Cm(0)
    if left is not None:
        pf.left_indent = Cm(left)
    mapping = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }
    pf.alignment = mapping[align]


def _shade(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_cell_border(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "666666")
        tc_borders.append(element)
    tc_pr.append(tc_borders)


def _prevent_row_split(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def _mark_header_row(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    tr_pr.append(header)


def _page_field(paragraph):
    run = paragraph.add_run()
    _set_run_font(run, FONT, 12)
    r = run._r

    def fld(kind):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        return element

    begin = fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = fld("separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = fld("end")
    r.append(begin)
    r2 = paragraph.add_run()._r
    r2.append(instr)
    r3 = paragraph.add_run()._r
    r3.append(separate)
    r4 = paragraph.add_run()
    _set_run_font(r4, FONT, 12)
    r4._r.append(text)
    r5 = paragraph.add_run()._r
    r5.append(end)


def _toc_field(paragraph):
    run = paragraph.add_run()
    r = run._r

    def fld(kind):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        return element

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    r.append(fld("begin"))
    r.append(instr)
    r.append(fld("separate"))
    placeholder = paragraph.add_run(
        "Оглавление формируется полем Word по заголовкам уровней 1–3. "
        "При открытии в Microsoft Word подтвердите обновление полей. "
        "Ниже приведён состав разделов, чтобы структура была видна и до обновления поля."
    )
    _set_run_font(placeholder, FONT, 14, italic=True)
    end_run = paragraph.add_run()
    end_run._r.append(fld("end"))


def _enable_update_fields(document):
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _restart_numbering(section):
    sect_pr = section._sectPr
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sect_pr.append(pg)


class Book:
    def __init__(self):
        self.document = Document()
        self.table_no = 0
        self.figure_no = 0
        self.listing_no = 0
        self._seen_h1 = False
        self._setup()

    def _setup(self):
        section = self.document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.header_distance = Mm(10)
        section.footer_distance = Mm(10)
        section.different_first_page_header_footer = True

        normal = self.document.styles["Normal"]
        _style_font(normal, FONT, 14)
        _paragraph_format(normal.paragraph_format, before=0, after=6, line=1.5, first=1.25)

        for style_name, size, bold, before, after, break_before in (
            ("Heading 1", 16, True, 0, 12, True),
            ("Heading 2", 14, True, 16, 8, False),
            ("Heading 3", 14, True, 12, 6, False),
        ):
            style = self.document.styles[style_name]
            _style_font(style, FONT, size, bold=bold)
            _paragraph_format(style.paragraph_format, before=before, after=after, line=1.15, align="left")
            style.paragraph_format.page_break_before = break_before
            style.paragraph_format.keep_with_next = True

        _enable_update_fields(self.document)

    def _p(self, text="", *, style=None, align="justify", first=1.25, before=0, after=6, size=14, bold=False, italic=False, center=False):
        paragraph = self.document.add_paragraph(style=style)
        if style is None:
            _paragraph_format(
                paragraph.paragraph_format,
                before=before,
                after=after,
                line=1.5,
                first=0 if center or align == "left" else first,
                align="center" if center else align,
            )
        if text:
            run = paragraph.add_run(text)
            _set_run_font(run, FONT, size, bold=bold, italic=italic)
        return paragraph

    def add_runs(self, paragraph, text, size=14, bold=False, italic=False, mono=False):
        run = paragraph.add_run(text)
        _set_run_font(run, MONO if mono else FONT, size, bold=bold, italic=italic)
        return run

    def title_page(self):
        self._p("Министерство образования Иркутской области", center=True, first=0, after=2, size=14)
        self._p(
            "Государственное бюджетное профессиональное образовательное учреждение Иркутской области",
            center=True,
            first=0,
            after=2,
            size=14,
        )
        self._p("«Иркутский авиационный техникум»", center=True, first=0, after=18, size=14, bold=True)
        self._p("УЧЕБНОЕ ПОСОБИЕ", center=True, first=0, before=24, after=8, size=16, bold=True)
        self._p(
            "Обеспечение безопасности веб-приложений:",
            center=True,
            first=0,
            after=2,
            size=16,
            bold=True,
        )
        self._p(
            "практикум по выявлению уязвимостей и разработке защитных механизмов",
            center=True,
            first=0,
            after=16,
            size=16,
            bold=True,
        )
        self._p(
            "Специальность 09.02.07 «Информационные системы и программирование»",
            center=True,
            first=0,
            after=4,
            size=14,
        )
        self._p(
            "ПМ.09 «Проектирование, разработка и оптимизация веб-приложений»",
            center=True,
            first=0,
            after=4,
            size=14,
        )
        self._p(
            "МДК.09.03 «Обеспечение безопасности веб-приложений»",
            center=True,
            first=0,
            after=18,
            size=14,
        )
        self._p("Рукопись для внутреннего рецензирования", center=True, first=0, after=18, size=14, italic=True)
        self._p("Автор: Бодоев Даниил Александрович", center=True, first=0, before=36, after=4, size=14)
        self._p("Иркутск, 2026", center=True, first=0, after=0, size=14)

    def start_body(self):
        new_section = self.document.add_section()
        new_section.page_width = Mm(210)
        new_section.page_height = Mm(297)
        new_section.left_margin = Mm(30)
        new_section.right_margin = Mm(15)
        new_section.top_margin = Mm(20)
        new_section.bottom_margin = Mm(20)
        new_section.header_distance = Mm(8)
        new_section.footer_distance = Mm(8)
        new_section.different_first_page_header_footer = False
        _restart_numbering(new_section)

        header = new_section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run("Безопасность веб-приложений: практикум")
        _set_run_font(run, FONT, 10, italic=True, color=(80, 80, 80))

        footer = new_section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _page_field(fp)

    def h(self, level, text):
        paragraph = self.document.add_heading(text, level=level)
        if level == 1:
            if not self._seen_h1:
                paragraph.paragraph_format.page_break_before = False
                self._seen_h1 = True
        return paragraph

    def para(self, text):
        self._p(text)

    def note(self, text):
        paragraph = self._p("", align="left", first=0, before=6, after=8)
        self.add_runs(paragraph, "Примечание. ", size=14, bold=True, italic=True)
        self.add_runs(paragraph, text, size=14, italic=True)

    def bullets(self, items):
        for item in items:
            paragraph = self.document.add_paragraph()
            _paragraph_format(paragraph.paragraph_format, before=0, after=2, line=1.15, first=0, left=1.0, align="left")
            self.add_runs(paragraph, "• " + item, size=14)

    def numbers(self, items):
        for index, item in enumerate(items, start=1):
            paragraph = self.document.add_paragraph()
            _paragraph_format(paragraph.paragraph_format, before=0, after=2, line=1.15, first=0, left=1.0, align="left")
            self.add_runs(paragraph, f"{index}. {item}", size=14)

    def caption(self, kind, text):
        if kind == "table":
            self.table_no += 1
            label = f"Таблица {self.table_no}. {text}"
        elif kind == "figure":
            self.figure_no += 1
            label = f"Рисунок {self.figure_no}. {text}"
        else:
            self.listing_no += 1
            label = f"Листинг {self.listing_no}. {text}"
        self._p(label, center=True, first=0, before=8, after=4, size=12, italic=True)

    def table(self, headers, rows, title):
        self.caption("table", title)
        table = self.document.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        hdr = table.rows[0].cells
        _mark_header_row(table.rows[0])
        _prevent_row_split(table.rows[0])
        for index, header in enumerate(headers):
            hdr[index].text = ""
            paragraph = hdr[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(header)
            _set_run_font(run, FONT, 11, bold=True)
            _shade(hdr[index], "E6E6E6")
            _set_cell_border(hdr[index])
        for row in rows:
            cells = table.add_row().cells
            _prevent_row_split(table.rows[-1])
            for index, value in enumerate(row):
                cells[index].text = ""
                paragraph = cells[index].paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = paragraph.add_run(str(value))
                _set_run_font(run, FONT, 11)
                _set_cell_border(cells[index])
        self._p("", first=0, after=6)

    def figure(self, title, art):
        self.caption("figure", title)
        for line in art.strip("\n").splitlines():
            paragraph = self.document.add_paragraph()
            _paragraph_format(paragraph.paragraph_format, before=0, after=0, line=1.0, first=0, align="left")
            self.add_runs(paragraph, line if line else " ", size=10, mono=True)
        self._p("", first=0, after=6)

    def code_block(self, title, code):
        self.caption("listing", title)
        for line in code.strip("\n").splitlines():
            paragraph = self.document.add_paragraph()
            _paragraph_format(paragraph.paragraph_format, before=0, after=0, line=1.0, first=0, left=0.2, align="left")
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F4F4F4")
            shading.set(qn("w:val"), "clear")
            paragraph._p.get_or_add_pPr().append(shading)
            self.add_runs(paragraph, line if line else " ", size=9, mono=True)
        self._p("", first=0, after=6)

    def code_file(self, relative, title):
        path = STAND / relative
        code = path.read_text(encoding="utf-8")
        self.code_block(title, code)

    def toc(self):
        self.h(1, "Содержание")
        paragraph = self.document.add_paragraph()
        _paragraph_format(paragraph.paragraph_format, before=0, after=8, line=1.15, first=0, align="left")
        _toc_field(paragraph)
        self.para(
            "Состав рукописи до обновления поля оглавления: сведения об издании; аннотация; целевая аудитория; "
            "назначение; связь с рабочей программой; список сокращений; введение; цель и задачи; планируемые результаты; "
            "глава 1 «Основы безопасности веб-приложений»; глава 2 «Методика аудита безопасности»; "
            "глава 3 «Практические и лабораторные работы»; глава 4 «Кейсы и задания для самостоятельной работы»; "
            "глава 5 «Контроль и оценивание»; заключение; источники; приложения А–Л."
        )

    def render(self, text):
        lines = text.splitlines()
        index = 0
        paragraph_buf = []

        def flush_paragraph():
            nonlocal paragraph_buf
            if paragraph_buf:
                self.para(" ".join(part.strip() for part in paragraph_buf))
                paragraph_buf = []

        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if stripped.startswith(":::"):
                flush_paragraph()
                head = stripped[3:].strip()
                kind, _, title = head.partition(" ")
                index += 1
                body = []
                while index < len(lines) and lines[index].strip() != ":::":
                    body.append(lines[index])
                    index += 1
                index += 1
                payload = "\n".join(body)
                if kind == "table":
                    rows = []
                    for raw in payload.splitlines():
                        if not raw.strip() or set(raw.strip()) <= {"-", "|", " "}:
                            continue
                        cells = [cell.strip() for cell in raw.strip().strip("|").split("|")]
                        rows.append(cells)
                    if not rows:
                        continue
                    self.table(rows[0], rows[1:], title)
                elif kind == "figure":
                    self.figure(title, payload)
                elif kind == "code":
                    self.code_block(title, payload)
                elif kind == "listing":
                    relative = title.strip()
                    caption = " ".join(payload.split()) or relative
                    self.code_file(relative, caption)
                elif kind == "note":
                    self.note(payload.strip())
                elif kind == "toc":
                    self.toc()
                elif kind == "bullets":
                    items = [item.strip()[2:] if item.strip().startswith("- ") else item.strip() for item in payload.splitlines() if item.strip()]
                    self.bullets(items)
                elif kind == "numbers":
                    items = []
                    for item in payload.splitlines():
                        item = item.strip()
                        if not item:
                            continue
                        items.append(re.sub(r"^\d+\.\s*", "", item))
                    self.numbers(items)
                else:
                    raise ValueError(f"Неизвестный блок {kind}")
                continue
            if stripped == "":
                flush_paragraph()
                index += 1
                continue
            if stripped.startswith("# "):
                flush_paragraph()
                self.h(1, stripped[2:].strip())
                index += 1
                continue
            if stripped.startswith("## "):
                flush_paragraph()
                self.h(2, stripped[3:].strip())
                index += 1
                continue
            if stripped.startswith("### "):
                flush_paragraph()
                self.h(3, stripped[4:].strip())
                index += 1
                continue
            if stripped.startswith("- "):
                flush_paragraph()
                items = []
                while index < len(lines) and lines[index].strip().startswith("- "):
                    items.append(lines[index].strip()[2:].strip())
                    index += 1
                self.bullets(items)
                continue
            if re.match(r"^\d+\.\s+", stripped):
                flush_paragraph()
                items = []
                while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                    items.append(re.sub(r"^\d+\.\s*", "", lines[index].strip()))
                    index += 1
                self.numbers(items)
                continue
            paragraph_buf.append(stripped)
            index += 1
        flush_paragraph()

    def save(self, path):
        self.document.save(path)


# Исправление: kind == "file" был черновиком. Используем только kind == "listing".
